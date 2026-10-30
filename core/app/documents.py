"""Documents JARVIS makes for the owner (create_document): Word, Excel, PDF, Markdown.

The writing itself is done by the model of the 'produktivitas' category (inside Hermes, or the main
agent); this module only turns structured content into a file, stores it under data/documents (backed
up) and lets the phone open it (/api/v1/documents/{id}).

Content format (simple Markdown): '# ' / '## ' / '### ' headings, '- ' bullets, '1. ' numbered items,
'| a | b |' tables (first row = header), '**bold**' inline, blank line = new paragraph. For xlsx, every
table becomes a sheet (or `rows` — a list of lists — is used directly).
"""

import os
import re
import uuid
from pathlib import Path

from . import audit, db

ROOT = Path(os.getenv("DOCUMENTS_DIR", "/data/documents"))
FORMATS = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
           "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
           "pdf": "application/pdf", "md": "text/markdown; charset=utf-8",
           "zip": "application/zip", "png": "image/png", "jpg": "image/jpeg", "txt": "text/plain; charset=utf-8"}   # the last four: files the tools made (store)
MAX_CONTENT = 400_000


class DocumentError(ValueError):
    pass


def parse(content: str) -> list[tuple]:
    """Markdown-ish text → blocks: ("h", level, text) | ("p", text) | ("ul", [items]) | ("ol", [items]) | ("table", rows)."""
    blocks: list[tuple] = []
    para: list[str] = []

    def flush():
        if para:
            blocks.append(("p", " ".join(para).strip()))
            para.clear()

    for raw in content.splitlines():
        line = raw.rstrip()
        if not line.strip():
            flush()
            continue
        if m := re.match(r"^(#{1,3})\s+(.*)", line):
            flush()
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
        elif m := re.match(r"^\s*[-*•]\s+(.*)", line):
            flush()
            if blocks and blocks[-1][0] == "ul":
                blocks[-1][1].append(m.group(1))
            else:
                blocks.append(("ul", [m.group(1)]))
        elif m := re.match(r"^\s*\d+[.)]\s+(.*)", line):
            flush()
            if blocks and blocks[-1][0] == "ol":
                blocks[-1][1].append(m.group(1))
            else:
                blocks.append(("ol", [m.group(1)]))
        elif line.strip().startswith("|"):
            flush()
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue  # header separator
            if blocks and blocks[-1][0] == "table":
                blocks[-1][1].append(cells)
            else:
                blocks.append(("table", [cells]))
        else:
            para.append(line.strip())
    flush()
    return blocks


def _runs(paragraph, text: str) -> None:
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = bool(i % 2)


def _docx(title: str, blocks: list[tuple], path: Path) -> None:
    from docx import Document
    doc = Document()
    doc.add_heading(title, 0)
    for b in blocks:
        if b[0] == "h":
            doc.add_heading(b[2], b[1])
        elif b[0] == "p":
            _runs(doc.add_paragraph(), b[1])
        elif b[0] in ("ul", "ol"):
            for item in b[1]:
                _runs(doc.add_paragraph(style="List Bullet" if b[0] == "ul" else "List Number"), item)
        elif b[0] == "table":
            rows = b[1]
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for i, row in enumerate(rows):
                for j in range(width):
                    cell = table.cell(i, j)
                    cell.text = row[j] if j < len(row) else ""
                    if i == 0:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                run.bold = True
    doc.save(path)


def _xlsx(title: str, blocks: list[tuple], rows: list | None, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    tables = [rows] if rows else [b[1] for b in blocks if b[0] == "table"]
    if not tables:
        tables = [[[b[-1] if not isinstance(b[-1], list) else "; ".join(b[-1])] for b in blocks]]
    for n, table in enumerate(tables, 1):
        ws = wb.create_sheet((title[:25] if n == 1 else f"Tabel {n}") or f"Tabel {n}")
        for r in table:
            ws.append([_number(c) for c in r])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) + 2 for c in col)))
    wb.save(path)


def _number(v):
    if isinstance(v, str):
        s = v.replace(".", "").replace(",", ".") if re.fullmatch(r"-?\d{1,3}(\.\d{3})+(,\d+)?", v) else v
        try:
            return float(s) if "." in s else int(s)
        except ValueError:
            return v
    return v


def _pdf(title: str, blocks: list[tuple], path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import ListFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    styles = getSampleStyleSheet()

    def rich(t: str) -> str:
        return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", escape(t))

    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 8)]
    for b in blocks:
        if b[0] == "h":
            story.append(Paragraph(rich(b[2]), styles[f"Heading{min(b[1] + 1, 4)}"]))
        elif b[0] == "p":
            story.append(Paragraph(rich(b[1]), styles["BodyText"]))
        elif b[0] in ("ul", "ol"):
            story.append(ListFlowable([Paragraph(rich(i), styles["BodyText"]) for i in b[1]],
                                      bulletType="bullet" if b[0] == "ul" else "1"))
        elif b[0] == "table":
            width = max(len(r) for r in b[1])
            data = [[Paragraph(rich(r[j] if j < len(r) else ""), styles["BodyText"]) for j in range(width)] for r in b[1]]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7"))]))
            story.append(t)
        story.append(Spacer(1, 6))
    SimpleDocTemplate(str(path), pagesize=A4, title=title).build(story)


async def create(title: str, fmt: str, content: str, actor: str, task_id: int | None = None,
                 rows: list | None = None) -> dict:
    title = (title or "Dokumen").strip()[:150]
    if fmt not in ("docx", "xlsx", "pdf", "md"):
        raise DocumentError("format harus docx, xlsx, pdf, atau md")
    if len(content or "") > MAX_CONTENT:
        raise DocumentError("isi dokumen terlalu panjang")
    if rows is not None and (not isinstance(rows, list) or not all(isinstance(r, list) for r in rows) or len(rows) > 20000):
        raise DocumentError("rows harus list of list (≤20.000 baris)")
    ROOT.mkdir(parents=True, exist_ok=True)
    doc_id = uuid.uuid4()
    safe = re.sub(r"[^A-Za-z0-9 _-]", "", title).strip().replace(" ", "_")[:60] or "dokumen"
    filename = f"{doc_id}.{fmt}"
    path = ROOT / filename
    blocks = parse(content or "")
    if fmt == "docx":
        _docx(title, blocks, path)
    elif fmt == "xlsx":
        _xlsx(title, blocks, rows, path)
    elif fmt == "pdf":
        _pdf(title, blocks, path)
    else:
        path.write_text(f"# {title}\n\n{content}", encoding="utf-8")
    row = await db.pool().fetchrow(
        """INSERT INTO documents (id, title, format, filename, size, task_id, created_by)
           VALUES ($1, $2, $3, $4, $5, $6, $7) RETURNING id, title, format, size, created_at""",
        doc_id, title, fmt, filename, path.stat().st_size, task_id, actor)
    await audit.record(actor, "document.created", f"document:{doc_id}", title=title, format=fmt)
    from . import canvas
    company = await db.pool().fetchval("SELECT company_id FROM tasks WHERE id = $1", task_id) if task_id else None
    await canvas.auto(f"doc:{doc_id}", company)
    return dict(row) | {"url": f"/api/v1/documents/{doc_id}", "download_name": f"{safe}.{fmt}"}


async def store(title: str, fmt: str, data: bytes, actor: str) -> dict:
    """A finished file a tool made (a merged PDF, page images in a ZIP, a Word copy …): kept and served like any document."""
    if fmt not in FORMATS:
        raise DocumentError("format berkas tidak dikenal")
    title = (title or "Berkas").strip()[:150]
    ROOT.mkdir(parents=True, exist_ok=True)
    doc_id = uuid.uuid4()
    filename = f"{doc_id}.{fmt}"
    (ROOT / filename).write_bytes(data)
    row = await db.pool().fetchrow(
        "INSERT INTO documents (id, title, format, filename, size, created_by) VALUES ($1, $2, $3, $4, $5, $6) RETURNING id, title, format, size, created_at",
        doc_id, title, fmt, filename, len(data), actor)
    await audit.record(actor, "document.stored", f"document:{doc_id}", title=title, format=fmt)
    safe = re.sub(r"[^A-Za-z0-9 _-]", "", title).strip().replace(" ", "_")[:60] or "berkas"
    return dict(row) | {"url": f"/api/v1/documents/{doc_id}", "download_name": f"{safe}.{fmt}"}


async def get(doc_id: str) -> tuple[dict, Path] | None:
    try:
        uid = uuid.UUID(doc_id)
    except ValueError:
        return None
    row = await db.pool().fetchrow("SELECT * FROM documents WHERE id = $1", uid)
    if not row:
        return None
    path = ROOT / row["filename"]
    return (dict(row), path) if path.is_file() else None


async def list_(limit: int = 50) -> list[dict]:
    rows = await db.pool().fetch(
        "SELECT id, title, format, size, task_id, created_by, created_at FROM documents ORDER BY created_at DESC LIMIT $1", limit)
    return [dict(r) | {"url": f"/api/v1/documents/{r['id']}"} for r in rows]
